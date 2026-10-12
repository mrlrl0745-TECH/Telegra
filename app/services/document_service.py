from io import BytesIO
from html import escape
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.schemas import LessonContent


BLUE = "00A0DF"
GRID = "A6A6A6"


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(table, color: str = GRID, size: str = "5") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _text(cell, value: str, *, bold: bool = False, center: bool = False, size: int = 9) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(value or "")
    run.bold = bold
    DocumentService._set_run_font(run, "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


class DocumentService:
    """Renders the uploaded reference form into an editable Word document."""

    @staticmethod
    def additional_sections(content: LessonContent) -> list[tuple[str, str]]:
        kk = content.language == "kk"
        return [
            ("Тілдік мақсаттар" if kk else "Языковые цели", content.language_objectives),
            ("Саралау" if kk else "Дифференциация", content.differentiation_plan),
            ("Бағалау" if kk else "Оценивание", content.assessment_plan),
            ("Пәнаралық байланыстар" if kk else "Межпредметные связи", content.cross_curricular_links),
            ("Денсаулық және қауіпсіздік" if kk else "Здоровье и безопасность", content.health_and_safety),
        ]

    @staticmethod
    def to_docx(data: LessonContent | dict) -> bytes:
        content = data if isinstance(data, LessonContent) else LessonContent.model_validate(data)
        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11.7), Inches(8.3)
        section.left_margin = section.right_margin = Inches(0.35)
        section.top_margin = section.bottom_margin = Inches(0.35)
        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        normal.font.size = Pt(9)

        title = doc.add_paragraph(style="Title")
        title_style = doc.styles["Title"]
        title_style.font.color.rgb = RGBColor(0, 0, 0)
        for properties in (title_style.element.pPr, title._p.pPr):
            if properties is not None:
                for border in list(properties.findall(qn("w:pBdr"))):
                    properties.remove(border)
        title.paragraph_format.space_after = Pt(4)
        title_run = title.add_run(content.title)
        DocumentService._set_run_font(title_run, "Arial")
        title_run.font.size = Pt(12)
        title_run.font.color.rgb = RGBColor(0, 0, 0)
        topic = doc.add_paragraph()
        topic.paragraph_format.space_after = Pt(8)
        topic_label = "Сабақ тақырыбы: " if content.language == "kk" else "Тема урока: "
        topic_label_run = topic.add_run(topic_label)
        topic_label_run.bold = True
        DocumentService._set_run_font(topic_label_run, "Arial")
        topic_label_run.font.size = Pt(10)
        topic_run = topic.add_run(content.lesson_topic)
        DocumentService._set_run_font(topic_run, "Arial")
        topic_run.font.size = Pt(10)

        school = doc.add_paragraph()
        school.add_run(("Мектеп / ұйым: " if content.language == "kk" else "Школа / организация: ") + (content.school or "____________________"))
        school.add_run(("    Пән: " if content.language == "kk" else "    Предмет: ") + content.subject)

        meta = doc.add_table(rows=7, cols=3)
        meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta.autofit = False
        widths = [Inches(3.5), Inches(3.75), Inches(3.75)]
        for column, width in zip(meta.columns, widths):
            column.width = width
        for row in meta.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = width
        _text(meta.cell(0, 0).merge(meta.cell(0, 2)), f"Бөлім: {content.section}" if content.language == "kk" else f"Раздел: {content.section}", size=10)
        _shade(meta.cell(0, 0), BLUE)
        _text(meta.cell(1, 0), "Педагогтің аты-жөні:" if content.language == "kk" else "ФИО педагога:", size=9)
        _text(meta.cell(1, 1).merge(meta.cell(1, 2)), content.teacher_name)
        _text(meta.cell(2, 0), "Күні:" if content.language == "kk" else "Дата:", size=9)
        _text(meta.cell(2, 1).merge(meta.cell(2, 2)), content.date)
        _text(meta.cell(3, 0), f"Сынып: {content.class_name}" if content.language == "kk" else f"Класс: {content.class_name}", size=9)
        _text(meta.cell(3, 1), f"Қатысқандар саны: {content.present_count}" if content.language == "kk" else f"Количество присутствующих: {content.present_count}")
        _text(meta.cell(3, 2), f"Қатыспағандар саны: {content.absent_count}" if content.language == "kk" else f"Количество отсутствующих: {content.absent_count}")
        _text(meta.cell(4, 0), "Оқу бағдарламасына сәйкес оқу мақсаттары" if content.language == "kk" else "Цели обучения в соответствии с учебной программой", size=9)
        _text(meta.cell(4, 1).merge(meta.cell(4, 2)), "\n".join(f"• {x}" for x in content.learning_objectives))
        _text(meta.cell(5, 0), "Сабақ мақсаттары" if content.language == "kk" else "Цели урока", size=9)
        _text(meta.cell(5, 1).merge(meta.cell(5, 2)), "\n".join(f"• {x}" for x in content.lesson_objectives))
        _text(meta.cell(6, 0).merge(meta.cell(6, 2)), "Сабақ барысы" if content.language == "kk" else "Ход урока", center=True, bold=True, size=11)
        _borders(meta)

        headings = ["Сабақ кезеңі/уақыты", "Педагогтің әрекеті", "Оқушының әрекеті", "Ресурстар", "Бағалау"] if content.language == "kk" else ["Этап урока/время", "Действия педагога", "Действия ученика", "Ресурсы", "Оценивание"]
        widths = [Inches(1.65), Inches(3.3), Inches(2.75), Inches(1.55), Inches(1.75)]
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for index, label in enumerate(headings):
            table.columns[index].width = widths[index]
            table.rows[0].cells[index].width = widths[index]
            _text(table.rows[0].cells[index], label, size=9)
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat_header = OxmlElement("w:tblHeader")
        repeat_header.set(qn("w:val"), "true")
        header_properties.append(repeat_header)
        for stage in content.stages:
            cells = table.add_row().cells
            for index, (value, width) in enumerate(zip([
                f"{stage.stage_name}\n{stage.time}", stage.teacher_actions, stage.student_actions, stage.resources, stage.assessment,
            ], widths)):
                cells[index].width = width
                _text(cells[index], value, size=8)
        _borders(table)

        extra = doc.add_table(rows=0, cols=2)
        extra.alignment = WD_TABLE_ALIGNMENT.CENTER
        extra.autofit = False
        extra.columns[0].width = Inches(2.6)
        extra.columns[1].width = Inches(8.4)
        for label, value in DocumentService.additional_sections(content):
            cells = extra.add_row().cells
            cells[0].width, cells[1].width = Inches(2.6), Inches(8.4)
            _text(cells[0], label, bold=True)
            _text(cells[1], value)
            for cell in cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
        _borders(extra)

        if content.homework:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_with_next = bool(content.reflection)
            paragraph.paragraph_format.space_before = Pt(5)
            run = paragraph.add_run("Үй тапсырмасы: " if content.language == "kk" else "Домашнее задание: ")
            run.bold = True
            DocumentService._set_run_font(run, "Arial")
            body_run = paragraph.add_run(content.homework)
            DocumentService._set_run_font(body_run, "Arial")
        if content.reflection:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(2)
            run = paragraph.add_run("Рефлексия: ")
            run.bold = True
            DocumentService._set_run_font(run, "Arial")
            reflection_run = paragraph.add_run(content.reflection)
            DocumentService._set_run_font(reflection_run, "Arial")
        else:
            # End the supplementary block when there are no trailing paragraphs.
            if not content.homework:
                for cell in extra.rows[-1].cells:
                    cell.paragraphs[-1].paragraph_format.keep_with_next = False
        output = BytesIO()
        doc.save(output)
        rendered = output.getvalue()
        DocumentService.validate_docx(rendered)
        return rendered

    @staticmethod
    def _set_run_font(run, font_name: str) -> None:
        run.font.name = font_name
        properties = run._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.append(fonts)
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{script}"), font_name)

    @staticmethod
    def validate_docx(data: bytes) -> None:
        """Fail closed before download if the in-memory OOXML package is incomplete."""
        if not data or len(data) < 4 or not data.startswith(b"PK\x03\x04"):
            raise ValueError("DOCX export is empty or has an invalid signature")
        try:
            with ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ValueError("DOCX export is missing required document parts")
                if archive.testzip() is not None:
                    raise ValueError("DOCX export contains a damaged package entry")
        except BadZipFile as exc:
            raise ValueError("DOCX export is not a valid ZIP package") from exc

    @staticmethod
    def to_pdf(data: LessonContent | dict) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        content = data if isinstance(data, LessonContent) else LessonContent.model_validate(data)
        font_regular, font_bold = DocumentService._register_pdf_fonts()
        styles = getSampleStyleSheet()
        body = ParagraphStyle("KspBody", parent=styles["Normal"], fontName=font_regular, fontSize=7.5, leading=9, spaceAfter=0)
        small = ParagraphStyle("KspSmall", parent=body, fontSize=7, leading=8.3)
        header = ParagraphStyle("KspHeader", parent=body, fontName=font_bold, fontSize=7.5, leading=9)
        centered = ParagraphStyle("KspCentered", parent=body, alignment=TA_CENTER, fontName=font_bold, fontSize=10)
        title_style = ParagraphStyle("KspTitle", parent=body, fontSize=11)

        def p(value: str, style=body) -> Paragraph:
            safe = escape(value or "").replace("\n", "<br/>")
            return Paragraph(safe, style)

        is_kk = content.language == "kk"
        output = BytesIO()
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=23, bottomMargin=23, title=content.title, author="КСП Генератор")
        topic_heading = ParagraphStyle("KspTopicHeading", parent=body, fontName=font_bold, fontSize=8, leading=10)
        story = [p(content.title, title_style), Spacer(1, 3), p(("Сабақ тақырыбы: " if is_kk else "Тема урока: ") + content.lesson_topic, topic_heading), Spacer(1, 6)]

        story.extend([p(("Мектеп / ұйым: " if is_kk else "Школа / организация: ") + (content.school or "____________________") + ("    Пән: " if is_kk else "    Предмет: ") + content.subject), Spacer(1, 5)])
        meta_rows = [
            [p(("Бөлім: " if is_kk else "Раздел: ") + content.section), "", ""],
            [p("Педагогтің аты-жөні:" if is_kk else "ФИО педагога:"), p(content.teacher_name), ""],
            [p("Күні:" if is_kk else "Дата:"), p(content.date), ""],
            [p(("Сынып: " if is_kk else "Класс: ") + content.class_name), p(("Қатысқандар саны: " if is_kk else "Количество присутствующих: ") + str(content.present_count), small), p(("Қатыспағандар саны: " if is_kk else "Количество отсутствующих: ") + str(content.absent_count), small)],
            [p("Оқу бағдарламасына сәйкес оқу мақсаттары" if is_kk else "Цели обучения в соответствии с учебной программой", small), p("\n".join(f"• {x}" for x in content.learning_objectives), small), ""],
            [p("Сабақ мақсаттары" if is_kk else "Цели урока"), p("\n".join(f"• {x}" for x in content.lesson_objectives), small), ""],
            [p("Сабақ барысы" if is_kk else "Ход урока", centered), "", ""],
        ]
        available_width = landscape(A4)[0] - 48
        meta_widths = [available_width * .36, available_width * .32, available_width * .32]
        meta_table = Table(meta_rows, colWidths=meta_widths, repeatRows=0, hAlign="LEFT")
        meta_commands = [
            ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#A6A6A6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("SPAN", (0, 0), (2, 0)), ("BACKGROUND", (0, 0), (2, 0), colors.HexColor("#00A0DF")),
            ("SPAN", (1, 1), (2, 1)), ("SPAN", (1, 2), (2, 2)),
            ("SPAN", (1, 4), (2, 4)), ("SPAN", (1, 5), (2, 5)),
            ("SPAN", (0, 6), (2, 6)),
        ]
        meta_table.setStyle(TableStyle(meta_commands))
        story.append(meta_table)
        story.append(Spacer(1, 0))

        headings = ["Сабақ кезеңі/уақыты", "Педагогтің әрекеті", "Оқушының әрекеті", "Ресурстар", "Бағалау"] if is_kk else ["Этап урока/время", "Действия педагога", "Действия ученика", "Ресурсы", "Оценивание"]
        rows = [[p(label, header) for label in headings]]
        for stage in content.stages:
            rows.append([
                p(f"{stage.stage_name}\n{stage.time}", small), p(stage.teacher_actions, small),
                p(stage.student_actions, small), p(stage.resources, small), p(stage.assessment, small),
            ])
        widths = [available_width * part for part in (.18, .29, .23, .14, .16)]
        lesson_table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        lesson_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#A6A6A6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(lesson_table)
        supplementary_start = len(story)
        extra = Table([[p(label, header), p(value)] for label, value in DocumentService.additional_sections(content)], colWidths=[available_width * .24, available_width * .76], hAlign="LEFT")
        extra.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#A6A6A6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.extend([Spacer(1, 5), extra])
        if content.homework:
            story.extend([Spacer(1, 6), p(("Үй тапсырмасы: " if is_kk else "Домашнее задание: ") + content.homework, body)])
        if content.reflection:
            story.extend([Spacer(1, 3), p(("Рефлексия: " if is_kk else "Рефлексия: ") + content.reflection, body)])
        story[supplementary_start:] = [KeepTogether(story[supplementary_start:])]
        document.build(story)
        rendered = output.getvalue()
        if not rendered.startswith(b"%PDF-"):
            raise ValueError("PDF export is empty or has an invalid signature")
        return rendered

    @staticmethod
    def _register_pdf_fonts() -> tuple[str, str]:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        regular_candidates = [
            Path("C:/Windows/Fonts/arial.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
        ]
        bold_candidates = [
            Path("C:/Windows/Fonts/arialbd.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
            Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf"),
        ]
        regular = next((path for path in regular_candidates if path.exists()), None)
        bold = next((path for path in bold_candidates if path.exists()), None)
        if regular:
            pdfmetrics.registerFont(TTFont("KspUnicode", str(regular)))
            if bold:
                pdfmetrics.registerFont(TTFont("KspUnicodeBold", str(bold)))
                return "KspUnicode", "KspUnicodeBold"
            return "KspUnicode", "KspUnicode"
        return "Helvetica", "Helvetica-Bold"
