from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models import Template, TemplateField, User

ALLOWED = {
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".png": {"image/png", "application/octet-stream"},
    ".jpg": {"image/jpeg", "application/octet-stream"},
    ".jpeg": {"image/jpeg", "application/octet-stream"},
}
MAX_TEMPLATE_TEXT = 12000
MAX_DOCX_ENTRIES = 2048
MAX_DOCX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_DOCX_ENTRY_BYTES = 16 * 1024 * 1024
MAX_PDF_PAGES = 100


class TemplateService:
    async def save_upload(self, db: AsyncSession, user: User, name: str, filename: str, content_type: str, content: bytes) -> Template:
        ext = Path(filename).suffix.lower()
        if ext not in ALLOWED or content_type not in ALLOWED[ext]:
            raise ValueError("Поддерживаются только DOCX, PDF, PNG и JPG")
        if len(content) > settings.max_upload_bytes:
            raise ValueError("Размер файла превышает допустимый лимит")
        signatures = {".pdf": content.startswith(b"%PDF-"), ".png": content.startswith(b"\x89PNG\r\n\x1a\n"), ".jpg": content.startswith(b"\xff\xd8\xff"), ".jpeg": content.startswith(b"\xff\xd8\xff"), ".docx": content.startswith(b"PK\x03\x04")}
        if not signatures[ext]:
            raise ValueError("Содержимое файла не соответствует его типу")
        structure: dict = {"document_type": "ksp", "sections": [], "fields": [], "tables": [], "analysis_status": "limited"}
        if ext == ".docx":
            structure = self._analyze_docx(content)
        elif ext == ".pdf":
            structure = self._analyze_pdf(content)
        else:
            structure["analysis_note"] = "Изображение сохранено. OCR-анализ будет добавлен отдельным этапом."
        user_dir = Path(settings.upload_dir) / user.id
        user_dir.mkdir(parents=True, exist_ok=True)
        target = user_dir / f"{uuid4().hex}{ext}"
        target.write_bytes(content)
        template = Template(user_id=user.id, name=name.strip()[:160] or Path(filename).stem, file_path=str(target), file_type=ext[1:], template_structure=structure)
        db.add(template)
        await db.flush()
        fields = structure.get("fields", [])
        for position, field in enumerate(fields):
            db.add(TemplateField(template_id=template.id, field_key=field.get("key", f"field_{position}"), label=field.get("label", ""), position=position))
        return template

    @staticmethod
    def _analyze_docx(content: bytes) -> dict:
        import io
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_ENTRIES:
                    raise ValueError("DOCX содержит слишком много элементов")
                if sum(item.file_size for item in entries) > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("Распакованный DOCX превышает допустимый размер")
                if any(item.file_size > MAX_DOCX_ENTRY_BYTES for item in entries):
                    raise ValueError("Элемент DOCX превышает допустимый размер")
                names = {item.filename for item in entries}
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ValueError("Некорректная структура DOCX")
        except BadZipFile as exc:
            raise ValueError("Повреждённый DOCX архив") from exc
        doc = Document(io.BytesIO(content))
        sections = [p.text.strip()[:300] for p in doc.paragraphs if p.text.strip()][:20]
        tables = []
        for table in doc.tables[:5]:
            rows = [[cell.text.strip()[:150] for cell in row.cells[:8]] for row in table.rows[:5]]
            tables.append({"columns": rows[0] if rows else [], "rows": len(rows)})
        return {"document_type": "ksp", "sections": sections[:40], "fields": [], "tables": tables, "analysis_status": "extracted"}

    @staticmethod
    def _analyze_pdf(content: bytes) -> dict:
        import io
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError("Зашифрованные PDF не поддерживаются")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF содержит слишком много страниц")
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:20])[:MAX_TEMPLATE_TEXT]
        sections = [line.strip() for line in text.splitlines() if line.strip()][:60]
        return {"document_type": "ksp", "sections": sections, "fields": [], "tables": [], "analysis_status": "text_extracted", "analysis_note": "PDF-схема извлечена из текстового слоя; сканированные страницы требуют OCR."}


async def get_owned_template(db: AsyncSession, user_id: str, template_id: str) -> Template | None:
    return await db.scalar(select(Template).where(Template.id == template_id, Template.user_id == user_id))
