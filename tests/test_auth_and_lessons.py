import asyncio
import hashlib
import hmac
from io import BytesIO
import json
import time
from urllib.parse import urlencode
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader
from sqlalchemy import select

from app.config import settings
from app.database import get_db
from app.main import app
from app.models import User
from app.security import validate_telegram_init_data
from app.services.openrouter_service import OpenRouterError
from app.schemas import LessonInput
from app.services.lesson_generator import LessonGenerator
from app.services.document_service import DocumentService
from app.services.ai_types import AIResult


def _signed_telegram_init_data(token: str, user_id: int, *, auth_date: int | None = None) -> str:
    timestamp = auth_date if auth_date is not None else int(time.time())
    fields = {"auth_date": str(timestamp), "user": json.dumps({"id": user_id, "first_name": "Teacher"})}
    check = "\n".join(f"{key}={value}" for key, value in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    return urlencode(fields)


def test_telegram_init_data_signature(monkeypatch):
    token = "test-bot-token"
    monkeypatch.setattr(settings, "telegram_bot_token", token)
    fields = {"auth_date": str(int(time.time())), "user": json.dumps({"id": 12345, "first_name": "Test"})}
    check = "\n".join(f"{key}={value}" for key, value in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    assert validate_telegram_init_data(urlencode(fields), token)["id"] == 12345
    fields["user"] = json.dumps({"id": 9})
    with pytest.raises(ValueError):
        validate_telegram_init_data(urlencode(fields), token)


@pytest.mark.asyncio
async def test_telegram_init_data_cannot_be_replayed(client, monkeypatch):
    token = "telegram-test-token"
    monkeypatch.setattr(settings, "telegram_bot_token", token)
    init_data = _signed_telegram_init_data(token, 712345)

    first = await client.post("/api/auth/telegram", json={"init_data": init_data})
    replay = await client.post("/api/auth/telegram", json={"init_data": init_data})

    assert first.status_code == 200
    assert replay.status_code == 401


@pytest.mark.asyncio
async def test_telegram_auth_ignores_frontend_user_id_and_rejects_old_init_data(client, monkeypatch):
    token = "telegram-test-token"
    monkeypatch.setattr(settings, "telegram_bot_token", token)
    valid = await client.post(
        "/api/auth/telegram",
        json={"init_data": _signed_telegram_init_data(token, 712345), "user_id": 999999},
    )
    init_data = _signed_telegram_init_data(token, 712345, auth_date=int(time.time()) - 901)
    response = await client.post("/api/auth/telegram", json={"init_data": init_data, "user_id": 999999})

    assert valid.status_code == 200
    assert valid.json()["user"]["telegram_user_id"] == 712345
    assert response.status_code == 401


@pytest.mark.asyncio
async def test_telegram_auth_endpoint_registers_three_free_generations(client, monkeypatch):
    token = "telegram-test-token"
    monkeypatch.setattr(settings, "telegram_bot_token", token)
    fields = {"auth_date": str(int(time.time())), "user": json.dumps({"id": 712345, "first_name": "Teacher", "username": "teacher"})}
    check = "\n".join(f"{key}={value}" for key, value in sorted(fields.items()))
    secret = hmac.new(b"WebAppData", token.encode(), hashlib.sha256).digest()
    fields["hash"] = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
    response = await client.post("/api/auth/telegram", json={"init_data": urlencode(fields)})
    assert response.status_code == 200
    assert response.json()["user"]["telegram_user_id"] == 712345
    assert response.json()["user"]["free_generations"] == 3


@pytest.mark.asyncio
async def test_ai_keeps_teacher_objectives_verbatim(lesson_payload):
    class Provider:
        async def generate_json(self, messages):
            return AIResult(data={
                "title": "КСП", "section": "Другой раздел", "teacher_name": "Другой учитель", "date": "01.01.2000",
                "subject": "Другой предмет", "class_name": "1", "present_count": 0, "absent_count": 0,
                "lesson_topic": "Другая тема", "learning_objectives": ["Перефразированная цель"],
                "lesson_objectives": ["Другая цель"], "stages": [
                    {"stage_name": "Вводная часть", "time": "5 минут", "teacher_actions": "Объясняет", "student_actions": "Слушают", "resources": "Доска", "assessment": "Устная обратная связь"},
                    {"stage_name": "Практика", "time": "40 минут", "teacher_actions": "Дает задание", "student_actions": "Решают", "resources": "Лист", "assessment": "Проверка по критериям"},
                ], "homework": "", "reflection": "",
            }, model="test-model")

    payload = {**lesson_payload, "language": "kk"}
    result = await LessonGenerator(Provider()).generate(LessonInput.model_validate(payload))
    assert result.data["learning_objectives"] == lesson_payload["learning_objectives"]
    assert result.data["lesson_objectives"] == lesson_payload["lesson_objectives"]
    assert result.data["teacher_name"] == lesson_payload["teacher_name"]
    document = Document(BytesIO(DocumentService.to_docx(result.data)))
    text = " ".join(paragraph.text for paragraph in document.paragraphs) + " " + " ".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "Педагогтің аты-жөні:" in text
    assert "Сабақ кезеңі/уақыты" in text


@pytest.mark.asyncio
async def test_registration_and_free_generation_limit(client, auth_headers, lesson_payload):
    usage = await client.get("/api/users/usage", headers=auth_headers)
    assert usage.json()["free_generations"] == 3
    for left in (2, 1, 0):
        response = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
        assert response.status_code == 200
        assert (await client.get("/api/users/usage", headers=auth_headers)).json()["free_generations"] == left
    blocked = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert blocked.status_code == 402


@pytest.mark.asyncio
async def test_regeneration_consumes_free_generation_and_cannot_bypass_quota(client, auth_headers, lesson_payload):
    created = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert created.status_code == 200

    regenerated = await client.post(f"/api/lessons/{created.json()['id']}/regenerate", headers=auth_headers)
    assert regenerated.status_code == 200
    assert (await client.get("/api/users/usage", headers=auth_headers)).json()["free_generations"] == 1

    second = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert second.status_code == 200

    bypass = await client.post(f"/api/lessons/{created.json()['id']}/regenerate", headers=auth_headers)
    assert bypass.status_code == 402


@pytest.mark.asyncio
async def test_parallel_generation_cannot_spend_the_same_free_credit_twice(client, auth_headers, lesson_payload, monkeypatch):
    async for db in app.dependency_overrides[get_db]():
        user = await db.scalar(select(User))
        assert user is not None
        user.free_generations = 1
        await db.commit()
        break

    entered_provider = asyncio.Event()
    release_provider = asyncio.Event()
    provider_calls = 0

    async def blocked_generation(self, lesson, template_structure=None):
        nonlocal provider_calls
        provider_calls += 1
        entered_provider.set()
        await release_provider.wait()
        return AIResult(data=LessonGenerator._sample(lesson), model="test-mode")

    monkeypatch.setattr(LessonGenerator, "generate", blocked_generation)
    first = asyncio.create_task(client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload))
    await asyncio.wait_for(entered_provider.wait(), timeout=3)
    second = asyncio.create_task(client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload))
    await asyncio.sleep(0.1)
    release_provider.set()
    responses = await asyncio.gather(first, second)

    assert sorted(response.status_code for response in responses) == [200, 402]
    assert provider_calls == 1


@pytest.mark.asyncio
async def test_lessons_and_templates_are_not_accessible_to_other_users(client, auth_headers, lesson_payload, monkeypatch):
    created = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert created.status_code == 200
    uploaded = await client.post(
        "/api/templates", headers=auth_headers, data={"name": "My template"},
        files={"file": ("safe.png", b"\x89PNG\r\n\x1a\nsmall image", "image/png")},
    )
    assert uploaded.status_code == 200

    token = "telegram-test-token"
    monkeypatch.setattr(settings, "telegram_bot_token", token)
    other_login = await client.post(
        "/api/auth/telegram", json={"init_data": _signed_telegram_init_data(token, 812345)},
    )
    assert other_login.status_code == 200
    other_headers = {"Authorization": f"Bearer {other_login.json()['access_token']}"}
    lesson_id = created.json()["id"]
    lesson_content = created.json()["content_json"]

    assert (await client.get(f"/api/lessons/{lesson_id}", headers=other_headers)).status_code == 404
    assert (await client.put(f"/api/lessons/{lesson_id}", headers=other_headers, json=lesson_content)).status_code == 404
    assert (await client.delete(f"/api/lessons/{lesson_id}", headers=other_headers)).status_code == 404
    assert (await client.post(f"/api/lessons/{lesson_id}/export/docx", headers=other_headers)).status_code == 404
    assert (await client.post(f"/api/lessons/{lesson_id}/copy", headers=other_headers)).status_code == 404
    assert (await client.post(f"/api/lessons/{lesson_id}/regenerate", headers=other_headers)).status_code == 404
    template_id = uploaded.json()["id"]
    assert (await client.get(f"/api/templates/{template_id}", headers=other_headers)).status_code == 404
    assert (await client.delete(f"/api/templates/{template_id}", headers=other_headers)).status_code == 404


@pytest.mark.asyncio
async def test_ai_error_does_not_consume_generation(client, auth_headers, lesson_payload, monkeypatch):
    from app.services.lesson_generator import LessonGenerator

    async def fail(*args, **kwargs):
        raise OpenRouterError("unavailable")

    monkeypatch.setattr(LessonGenerator, "generate", fail)
    response = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert response.status_code == 503
    assert (await client.get("/api/users/usage", headers=auth_headers)).json()["free_generations"] == 3


@pytest.mark.asyncio
async def test_export_rendering_error_does_not_consume_generation(client, auth_headers, lesson_payload, monkeypatch):
    from app.services.document_service import DocumentService

    def fail(*args, **kwargs):
        raise RuntimeError("renderer failed")

    monkeypatch.setattr(DocumentService, "to_pdf", fail)
    response = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert response.status_code == 422
    assert (await client.get("/api/users/usage", headers=auth_headers)).json()["free_generations"] == 3


@pytest.mark.asyncio
async def test_lesson_crud_copy_and_docx_export(client, auth_headers, lesson_payload):
    created = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert created.status_code == 200
    lesson = created.json()
    assert lesson["content_json"]["learning_objectives"] == lesson_payload["learning_objectives"]
    copied = await client.post(f"/api/lessons/{lesson['id']}/copy", headers=auth_headers)
    assert copied.status_code == 200
    content = lesson["content_json"]
    content["lesson_topic"] = "Линейные уравнения: практика"
    updated = await client.put(f"/api/lessons/{lesson['id']}", headers=auth_headers, json=content)
    assert updated.status_code == 200
    assert updated.json()["topic"] == content["lesson_topic"]
    exported = await client.post(f"/api/lessons/{lesson['id']}/export/docx", headers=auth_headers)
    assert exported.status_code == 200
    assert "wordprocessingml" in exported.headers["content-type"]
    document = Document(__import__("io").BytesIO(exported.content))
    assert len(document.tables) == 2
    assert len(document.tables[1].columns) == 5
    pdf = await client.post(f"/api/lessons/{lesson['id']}/export/pdf", headers=auth_headers)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF-")
    assert len(PdfReader(BytesIO(pdf.content)).pages) >= 1
    deleted = await client.delete(f"/api/lessons/{lesson['id']}", headers=auth_headers)
    assert deleted.status_code == 200


@pytest.mark.asyncio
async def test_complete_two_user_archive_copy_edit_and_download_isolation(client, auth_headers, lesson_payload, monkeypatch):
    created = await client.post("/api/lessons/generate", headers=auth_headers, json=lesson_payload)
    assert created.status_code == 200
    original = created.json()
    original_id = original["id"]

    exported = await client.post(f"/api/lessons/{original_id}/export/docx", headers=auth_headers)
    assert exported.status_code == 200
    assert exported.content and exported.headers["content-length"] == str(len(exported.content))
    assert ".." not in exported.headers["content-disposition"]
    with ZipFile(BytesIO(exported.content)) as package:
        assert package.testzip() is None
        assert "word/document.xml" in package.namelist()
    opened = Document(BytesIO(exported.content))
    assert len(opened.tables) == 2

    a_archive = await client.get("/api/lessons", headers=auth_headers)
    assert [item["id"] for item in a_archive.json()] == [original_id]
    copied = await client.post(f"/api/lessons/{original_id}/copy", headers=auth_headers)
    assert copied.status_code == 200
    copy_content = copied.json()["content_json"]
    copy_content["lesson_topic"] = "Линейные уравнения: новый класс"
    copy_content["class_name"] = "8Б"
    copy_content["date"] = "2026-10-04"
    copy_content["learning_objectives"] = ["8.2.2.1 решать линейные уравнения"]
    edited = await client.put(f"/api/lessons/{copied.json()['id']}", headers=auth_headers, json=copy_content)
    assert edited.status_code == 200
    assert edited.json()["grade"] == "8Б"
    assert edited.json()["lesson_date"] == "2026-10-04"
    assert edited.json()["content_json"]["learning_objectives"] == copy_content["learning_objectives"]

    monkeypatch.setattr(settings, "telegram_bot_token", "telegram-test-token")
    other_login = await client.post("/api/auth/telegram", json={"init_data": _signed_telegram_init_data("telegram-test-token", 912345)})
    assert other_login.status_code == 200
    b_headers = {"Authorization": f"Bearer {other_login.json()['access_token']}"}
    assert (await client.get("/api/lessons", headers=b_headers)).json() == []

    for lesson_id in (original_id, copied.json()["id"]):
        assert (await client.get(f"/api/lessons/{lesson_id}", headers=b_headers)).status_code == 404
        assert (await client.put(f"/api/lessons/{lesson_id}", headers=b_headers, json=copy_content)).status_code == 404
        assert (await client.delete(f"/api/lessons/{lesson_id}", headers=b_headers)).status_code == 404
        assert (await client.post(f"/api/lessons/{lesson_id}/copy", headers=b_headers)).status_code == 404
        assert (await client.post(f"/api/lessons/{lesson_id}/regenerate", headers=b_headers)).status_code == 404
        assert (await client.post(f"/api/lessons/{lesson_id}/export/docx", headers=b_headers)).status_code == 404
        assert (await client.post(f"/api/lessons/{lesson_id}/export/pdf", headers=b_headers)).status_code == 404

    assert (await client.get(f"/api/lessons/{original_id}", headers=auth_headers)).status_code == 200
    assert (await client.get(f"/api/lessons/{copied.json()['id']}", headers=auth_headers)).json()["topic"] == "Линейные уравнения: новый класс"
    assert len((await client.get("/api/lessons", headers=auth_headers)).json()) == 2


@pytest.mark.asyncio
async def test_substitute_teacher_can_generate_and_download_a_self_contained_lesson(client, auth_headers, monkeypatch):
    monkeypatch.setattr(settings, "ai_test_mode", False)
    observed: dict[str, object] = {}

    class Provider:
        async def generate_json(self, messages):
            observed["system"] = messages[0]["content"]
            observed["request"] = json.loads(messages[1]["content"])
            return AIResult(data={
                "title": "КСП", "language": "ru", "section": "Не указан", "teacher_name": "ignored", "date": "ignored",
                "subject": "Математика", "class_name": "5А", "present_count": 0, "absent_count": 0,
                "lesson_topic": "Доли и дроби", "lesson_duration": 40, "substitute_mode": True,
                "learning_objectives": ["Сравнивать дроби с одинаковыми знаменателями без кода"],
                "lesson_objectives": ["Сравнить две дроби и объяснить способ сравнения"],
                "stages": [
                    {"stage_name": "Старт", "time": "5 минут", "teacher_actions": "Покажите карточки с дробями 1/5 и 3/5.", "student_actions": "Назовите, какая дробь больше.", "resources": "Доска", "assessment": "Устный ответ"},
                    {"stage_name": "Практика", "time": "20 минут", "teacher_actions": "Дайте задание: сравнить 2/7 и 5/7. Ключ: 5/7 больше 2/7.", "student_actions": "Запишите знак сравнения и объясните ответ.", "resources": "Карточки", "assessment": "Критерий: верно выбран знак; дескриптор: указаны числители и объяснён общий знаменатель"},
                    {"stage_name": "Итог", "time": "15 минут", "teacher_actions": "Соберите ответы и проведите рефлексию.", "student_actions": "Завершите фразу: теперь я умею…", "resources": "Тетрадь", "assessment": "Самооценка по цели"},
                ],
                "homework": "Сравнить 3/8 и 6/8, записать объяснение.", "reflection": "Назвать один освоенный приём.",
            }, model="test-model")

    monkeypatch.setattr("app.services.lesson_generator.get_ai_provider", lambda: Provider())
    response = await client.post("/api/lessons/generate", headers=auth_headers, json={
        "subject": "Математика", "grade": "5А", "topic": "Доли и дроби", "learning_objectives": [],
        "lesson_objectives": [], "lesson_duration": 40, "substitute_mode": True,
    })
    assert response.status_code == 200, response.text
    lesson = response.json()
    content = lesson["content_json"]
    assert observed["request"]["substitute_mode"] is True
    assert "Режим учителя на замене" in observed["system"]
    assert content["substitute_mode"] is True
    assert sum(int(stage["time"].split()[0]) for stage in content["stages"]) == 40
    assert "Ключ: 5/7 больше 2/7" in content["stages"][1]["teacher_actions"]
    assert lesson["teacher_name"] == "Учитель на замене"
    archive = await client.get("/api/lessons", headers=auth_headers)
    assert [item["id"] for item in archive.json()] == [lesson["id"]]

    exported = await client.post(f"/api/lessons/{lesson['id']}/export/docx", headers=auth_headers)
    assert exported.status_code == 200
    rendered = Document(BytesIO(exported.content))
    document_text = " ".join(paragraph.text for paragraph in rendered.paragraphs) + " " + " ".join(cell.text for table in rendered.tables for row in table.rows for cell in row.cells)
    assert "Ключ: 5/7 больше 2/7" in document_text


@pytest.mark.asyncio
async def test_template_upload_and_ownership(client, auth_headers):
    response = await client.post("/api/templates", headers=auth_headers, data={"name": "Шаблон школы"}, files={"file": ("school.png", b"\x89PNG\r\n\x1a\nimage-bytes", "image/png")})
    assert response.status_code == 200
    assert response.json()["name"] == "Шаблон школы"
    assert response.json()["file_type"] == "png"
    word = Document()
    word.add_paragraph("Краткосрочный план")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Действия педагога"
    table.rows[0].cells[1].text = "Действия ученика"
    file = BytesIO()
    word.save(file)
    docx_response = await client.post("/api/templates", headers=auth_headers, data={"name": "Школьный DOCX"}, files={"file": ("school.docx", file.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert docx_response.status_code == 200
    assert docx_response.json()["template_structure"]["analysis_status"] == "extracted"
    assert docx_response.json()["template_structure"]["tables"][0]["columns"] == ["Действия педагога", "Действия ученика"]

