from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from app.schemas import LessonContent
from app.services.lesson_generator import lesson_output_schema


@pytest.mark.asyncio
async def test_merged_fields_survive_save_copy_and_both_exports(client, auth_headers, lesson_payload):
    response = await client.post('/api/lessons/generate', headers=auth_headers, json=lesson_payload)
    assert response.status_code == 200
    lesson = response.json()
    content = lesson['content_json']
    fields = dict(school='Школа Альфа', language_objectives='Термин дискриминант',
                  differentiation_plan='Подсказка к первому шагу', assessment_plan='Проверить оба корня',
                  cross_curricular_links='Связь с физикой', health_and_safety='Перерыв для глаз')
    content.update(fields)
    saved = await client.put(f"/api/lessons/{lesson['id']}", headers=auth_headers, json=content)
    assert saved.status_code == 200
    copied = await client.post(f"/api/lessons/{lesson['id']}/copy", headers=auth_headers)
    assert copied.status_code == 200
    for key, value in fields.items():
        assert copied.json()['content_json'][key] == value
    for fmt in ('docx', 'pdf'):
        exported = await client.post(f"/api/lessons/{lesson['id']}/export/{fmt}", headers=auth_headers)
        assert exported.status_code == 200
        if fmt == 'docx':
            doc = Document(BytesIO(exported.content))
            text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        else:
            text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(exported.content)).pages)
        for value in fields.values():
            assert value in text
    legacy = {k: v for k, v in content.items() if k not in fields}
    assert LessonContent.model_validate(legacy).school == ''
    schema = lesson_output_schema()
    assert 'school' not in schema['properties']
    assert set(fields) - {'school'} <= set(schema['required'])
